"""Turns the organisers' final data workbook into the three files the server loads.

    python tools/import_final.py Stochasticdata.xlsx apps/api/scenarios/final

Writes:
  universe.json  the companies (the game names, never the real ones) and their opening prices
  prices.json    the exact price of every company at every 10 seconds of open-market time
  scenario.json  the news schedule, and where the prices file is

Prices follow the table exactly. The news is released by itself at its time in the run sheet (measured in
open-market minutes, so pausing or changing the schedule cannot put it out of step with the prices), and the
organiser can hold, edit, move or release any item.

This output is the event's secret (future prices and news). It is git-ignored: do not commit it.
"""
import json, os, sys, re
import openpyxl

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)

def secs(t):
    m = re.match(r'T\+(\d+):(\d+):(\d+)', str(t))
    return int(m[1]) * 3600 + int(m[2]) * 60 + int(m[3])

# ---- prices
ws = wb['PRICES_10S']
rows = list(ws.iter_rows(values_only=True))
head = rows[1]
symbols = [c for c in head[5:] if c]
table, event_sec_to_bar = [], {}
for r in rows[2:]:
    if r[0] is None:
        continue
    bar = int(r[0])
    if bar != len(table):
        sys.exit("price bars are not consecutive at %s" % bar)
    event_sec_to_bar[int(r[1])] = bar
    table.append([round(float(v), 2) for v in r[5:5 + len(symbols)]])
bar_seconds = int(rows[3][1] - rows[2][1])

# ---- companies
uni = []
for r in wb['UNIVERSE'].iter_rows(min_row=3, values_only=True):
    if not r[0]:
        continue
    uni.append({"symbol": r[0], "name": r[1], "sector": r[3]})
if [u['symbol'] for u in uni] != symbols:
    sys.exit("the companies in UNIVERSE and PRICES_10S do not match")
for i, u in enumerate(uni):
    u['openPrice'] = table[0][i]

# ---- news
events = []
for r in wb['RUN_SHEET'].iter_rows(min_row=3, values_only=True):
    if not r[0]:
        continue
    n, typ, cat, headline = int(r[0]), r[4], r[5], str(r[8]).strip()
    headline = headline.replace(' — ', ': ').replace(' � ', ': ')
    t = secs(r[2])
    if t not in event_sec_to_bar:
        sys.exit("news %d at %s is not inside live trading" % (n, r[2]))
    events.append({"id": "n%02d" % n, "atMinute": round(event_sec_to_bar[t] * bar_seconds / 60, 4), "headline": headline,
                   "body": "", "type": typ, "category": cat, "impacts": []})
events.sort(key=lambda e: e['atMinute'])

scenario = {"seed": 1, "tickSeconds": bar_seconds, "newsClock": "market", "pricesFile": "prices.json",
            "volatility": {"defaultPct": 0}, "events": events, "regimes": []}

os.makedirs(out, exist_ok=True)
json.dump(uni, open(os.path.join(out, 'universe.json'), 'w', encoding='utf8'), indent=1, ensure_ascii=False)
json.dump({"barSeconds": bar_seconds, "symbols": symbols, "rows": table}, open(os.path.join(out, 'prices.json'), 'w', encoding='utf8'), separators=(',', ':'))
json.dump(scenario, open(os.path.join(out, 'scenario.json'), 'w', encoding='utf8'), indent=1, ensure_ascii=False)
print(len(uni), "companies,", len(table), "price steps of", bar_seconds, "s,", len(events), "news items")
