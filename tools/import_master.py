"""Turns the organisers' master workbook into the two files the server loads.

    python tools/import_master.py Stockastic_Master.xlsx apps/api/scenarios/mock

Writes universe.json (companies and opening prices) and scenario.json (volatility, price limits, market
events with their price impacts, and bull and bear runs). Run it again whenever the workbook changes.

How the workbook is read (change the constants below if the final data means something different):
  * UNIVERSE: ticker, sim_name, sector, P_sim (opening price), vol_band (per company).
  * RUN_SHEET: fund release time "T+HH:MM:SS" is time since the event started. FAKE and DENIAL rows are
    news that moves nothing. REGIME rows start a bull or bear run that lasts until the end of that block.
  * IMPACT_MAP: one row per sector or company hit, with the impact as a fraction (0.02 = 2%).
  * PARAMETERS: price floor, cap, rounding and the price impact coefficient.
"""
import json, re, sys, os
import openpyxl

VOL_BAND_TO_TICK_PCT = 0.1   # a vol_band of 0.02 (2%) becomes a 0.2% standard deviation per price change
RAMP_MINUTES = 3             # an event's price move is spread over this many minutes
REGIME_DRIFT = {"MODERATE": 0.05, "MAJOR": 0.10, "MINOR": 0.03}   # % per price change
REGIME_BREADTH = 0.7
TICK_SECONDS = 60

src, out = sys.argv[1], sys.argv[2]
rb = json.load(open(os.path.join(os.path.dirname(__file__), '..', 'apps/api/internal/rulebook/rulebook.json'), encoding='utf8'))
wb = openpyxl.load_workbook(src, data_only=True)

def secs(t):
    m = re.match(r'T\+(\d+):(\d+):(\d+)', str(t))
    return int(m[1]) * 3600 + int(m[2]) * 60 + int(m[3])

params = {r[1]: r[2] for r in wb['PARAMETERS'].iter_rows(min_row=5, values_only=True) if r[1]}
coef = float(params.get('Price impact coefficient', 1))

uni, vol = [], {}
for r in wb['UNIVERSE'].iter_rows(min_row=3, values_only=True):
    if not r[0]:
        continue
    uni.append({"symbol": r[0], "name": r[1], "sector": r[3], "openPrice": float(r[6])})
    vol[r[0]] = round(float(r[14]) * 100 * VOL_BAND_TO_TICK_PCT, 4)

# where each block's trading ends, in minutes since the event started
ends, t = {}, 0
names = {"p1_trading": "P1", "p2_t1": "B1", "p2_t2": "B2", "p2_t3": "B3", "p2_t4": "B4"}
for b in rb["event"]["timeline"]:
    t += b["durationMin"]
    if b["id"] in names:
        ends[names[b["id"]]] = t

impacts = {}
for r in wb['IMPACT_MAP'].iter_rows(min_row=3, values_only=True):
    if not r[0]:
        continue
    im = {"shockPct": round(float(r[7]) * 100 * coef, 4), "rampMinutes": RAMP_MINUTES}
    im["sector" if r[3] == "SECTOR" else "symbol"] = r[4]
    impacts.setdefault(int(r[0]), []).append(im)

events, regimes = [], []
for r in wb['RUN_SHEET'].iter_rows(min_row=3, values_only=True):
    if not r[0]:
        continue
    n, block, typ, cat, mag, headline = int(r[0]), r[1], r[4], r[5], r[6], r[7]
    at = round(secs(r[2]) / 60, 4)
    eid = "e%02d" % n
    if typ == "REGIME":
        kind = "bull" if headline.upper().startswith("BULL") else "bear"
        regimes.append({"id": eid, "kind": kind, "atMinute": at, "durationMinutes": round(max(ends[block] - at, 1), 4),
                        "driftPctPerTick": REGIME_DRIFT.get(mag, 0.05), "breadth": REGIME_BREADTH, "headline": headline, "body": ""})
    else:
        events.append({"id": eid, "atMinute": at, "headline": headline, "body": "", "type": typ, "category": cat,
                       "impacts": impacts.get(n, []) if typ == "REAL" else []})

sc = {"seed": 20260921, "tickSeconds": TICK_SECONDS, "volatility": {"defaultPct": 0.3, "bySymbol": vol},
      "priceFloor": float(params.get('Price floor (Rs)', 0)), "priceCap": float(params.get('Price cap (Rs)', 0)),
      "roundTo": float(params.get('Price rounding (Rs)', 0)), "events": events, "regimes": regimes}
os.makedirs(out, exist_ok=True)
json.dump(uni, open(os.path.join(out, 'universe.json'), 'w', encoding='utf8'), indent=1, ensure_ascii=False)
json.dump(sc, open(os.path.join(out, 'scenario.json'), 'w', encoding='utf8'), indent=1, ensure_ascii=False)
print(len(uni), "companies,", len(events), "events,", len(regimes), "regimes")
