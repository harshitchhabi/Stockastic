"""Checks that the generated data files agree with every sheet of the organisers' workbook.

    python tools/verify_final.py [workbook.xlsx] [folder with universe.json, scenario.json, prices.json]

Run it after tools/import_final.py, and again whenever the workbook changes.
"""
import openpyxl, json, re, collections
import sys, os
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
wb = openpyxl.load_workbook(sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'Stochasticdata.xlsx'), data_only=True)
F = (sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'apps/api/scenarios/final')) + '/'
uni = json.load(open(F + 'universe.json', encoding='utf8'))
sc = json.load(open(F + 'scenario.json', encoding='utf8'))
pt = json.load(open(F + 'prices.json', encoding='utf8'))
def secs(t):
    m = re.match(r'T\+(\d+):(\d+):(\d+)', str(t)); return int(m[1])*3600+int(m[2])*60+int(m[3])
res = []
def check(name, ok, detail=''):
    res.append((name, ok, detail)); print(('PASS ' if ok else 'FAIL ') + name + (' | ' + detail if detail else ''))

print('sheets:', wb.sheetnames)
# ---- UNIVERSE
U = [r for r in wb['UNIVERSE'].iter_rows(min_row=3, values_only=True) if r[0]]
check('universe: 150 companies loaded, same tickers and order', [u['symbol'] for u in uni] == [r[0] for r in U] and len(U) == 150)
check('universe: game names used, not real names', all(u['name'] == r[1] for u, r in zip(uni, U)) and not any(u['name'] == r[2] for u, r in zip(uni, U)))
check('universe: sectors match', all(u['sector'] == r[3] for u, r in zip(uni, U)))
check('universe: 15 sectors', len({u['sector'] for u in uni}) == 15)
# ---- SECTOR_ALLOCATION
SA = {r[0]: r[4] for r in wb['SECTOR_ALLOCATION'].iter_rows(min_row=5, values_only=True) if r[0] and r[0] != 'TOTAL'}
cnt = collections.Counter(u['sector'] for u in uni)
check('sector allocation: slots per sector match the loaded companies', all(cnt[k] == v for k, v in SA.items()), str({k: (cnt[k], v) for k, v in SA.items() if cnt[k] != v}))
# ---- PRICES
rows = [r for r in wb['PRICES_10S'].iter_rows(min_row=3, values_only=True) if r[0] is not None]
head = [c for c in list(wb['PRICES_10S'].iter_rows(min_row=2, max_row=2, values_only=True))[0][5:] if c]
check('prices: 1020 steps x 150 companies, same columns as universe', len(pt['rows']) == 1020 and pt['symbols'] == head == [u['symbol'] for u in uni])
maxdiff = max(abs(round(float(v), 2) - pt['rows'][i][j]) for i, r in enumerate(rows) for j, v in enumerate(r[5:5+150]))
check('prices: every cell equal to the sheet', maxdiff < 1e-9, f'max diff {maxdiff}')
check('prices: opening price = bar 0', all(u['openPrice'] == pt['rows'][0][i] for i, u in enumerate(uni)))
# ---- RUN_SHEET vs scenario
RS = [r for r in wb['RUN_SHEET'].iter_rows(min_row=3, values_only=True) if r[0]]
ev = {e['id']: e for e in sc['events']}
check('run sheet: 40 items loaded, ids n01..n40', len(RS) == 40 == len(ev))
bar = {int(r[1]): int(r[0]) for r in rows}
bad = [n for n in RS if abs(ev['n%02d' % n[0]]['atMinute'] - bar[secs(n[2])] * 10 / 60) > 1e-6]
check('run sheet: every item at the right open-market minute', not bad)
check('run sheet: types and categories kept (REAL/FAKE/DENIAL/REGIME)', all(ev['n%02d' % n[0]]['type'] == n[4] and ev['n%02d' % n[0]]['category'] == n[5] for n in RS))
lag_ok = all((secs(n[3]) - secs(n[2])) == (0 if n[1] == 'P1' else 60) for n in RS)
check('run sheet: public release is +0 s in Phase 1 and +60 s in Phase 2 (server: 60 s lead in Phase 2 only)', lag_ok)
# ---- FEEDS
FF = [r for r in wb['FUND_FEED'].iter_rows(min_row=3, values_only=True) if r[0]]
PF = [r for r in wb['PARTICIPANT_FEED'].iter_rows(min_row=3, values_only=True) if r[0]]
same_text = all(str(f[3]).strip() == str(n[8]).strip() and str(p[3]).strip() == str(n[8]).strip() for f, p, n in zip(FF, PF, RS))
check('feeds: fund feed and participant feed text identical to the run sheet', same_text and len(FF) == len(PF) == 40)
check('feeds: fund feed times = run sheet fund release; participant feed times = public release',
      all(secs(f[1]) == secs(n[2]) and secs(p[1]) == secs(n[3]) for f, p, n in zip(FF, PF, RS)))
# ---- IMPACT_MAP: does the price table really move the named targets at the public release?
IM = [r for r in wb['IMPACT_MAP'].iter_rows(min_row=3, values_only=True) if r[0]]
sym_i = {s: i for i, s in enumerate(pt['symbols'])}
sec_members = collections.defaultdict(list)
for i, u in enumerate(uni): sec_members[u['sector']].append(i)
wrong = []
checked = 0
for r in IM:
    n, pub, tt, tgt, imp = r[0], secs(r[1]), r[4], r[5], float(r[8])
    if r[3] == 'REGIME':
        continue
    b = bar.get(pub)
    if b is None: wrong.append((n, 'no bar')); continue
    idx = sec_members[tgt] if tt == 'SECTOR' else [sym_i[tgt]]
    # average move over the 3 minutes after the public release vs the price just before it
    a = sum(pt['rows'][b - 1][i] for i in idx) / len(idx)
    z = sum(pt['rows'][min(b + 18, 1019)][i] for i in idx) / len(idx)
    move = (z - a) / a
    checked += 1
    if (move > 0) != (imp > 0) or abs(move) < 0.2 * abs(imp):
        wrong.append((n, tgt, imp, round(move, 4)))
check('impact map: the price table moves each target in the right direction after the public release', not wrong, f'{checked} rows checked, {len(wrong)} disagree: {wrong[:8]}')
# ---- RUMOUR_SPIKES
RU = [r for r in wb['RUMOUR_SPIKES'].iter_rows(min_row=3, values_only=True) if r[0] is not None and r[2]]
check('rumour spikes: 4 pairs, every rumour has a later denial', len({r[0] for r in RU}) == 4 and all(any(x[0] == r[0] and x[2] == 'DENIAL' and secs(x[3]) > secs(r[3]) for x in RU) for r in RU if r[2] == 'RUMOUR'))
# headline texts of rumours/denials exist as news items
heads = {e['headline'].strip() for e in sc['events']}
check('rumour spikes: each rumour and denial headline is in the news schedule', all(str(r[12]).strip() in heads for r in RU), str([r[12][:40] for r in RU if str(r[12]).strip() not in heads][:3]))
# ---- PARAMETERS
P = {r[1]: r[2] for r in wb['PARAMETERS'].iter_rows(min_row=5, values_only=True) if r[1]}
rb = json.load(open(os.path.join(ROOT, 'apps/api/internal/rulebook/rulebook.json'), encoding='utf8'))
check('parameters: starting wallet = rulebook starting capital', P['Starting wallet (Rs)'] == rb['accounts']['startingCapital'])
check('parameters: 25% single-stock limit in the rulebook', P['Max single-stock holding'] == rb['market']['maxSingleStockPercent'] / 100)
check('parameters: 60 s fund news lead in the rulebook', P['Fund news lead (s)'] == rb['news']['fundManagerLeadSeconds'])
check('parameters: price bar 10 s = tickSeconds', P['Price bar (s)'] == sc['tickSeconds'] == pt['barSeconds'])
# ---- DATA_CHECKS
DC = [r for r in wb['DATA_CHECKS'].iter_rows(min_row=3, values_only=True) if r[0]]
check('data checks sheet: no FAIL', not any(r[1] == 'FAIL' for r in DC), ', '.join(f'{r[0]}={r[1]}' for r in DC if r[1] not in ('PASS',)))
print('\n%d of %d checks passed' % (sum(1 for r in res if r[1]), len(res)))
